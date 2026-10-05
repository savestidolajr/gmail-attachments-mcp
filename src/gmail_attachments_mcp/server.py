"""MCP server: list, download and read Gmail attachments. Read-only."""
import base64
import io
import os
import re
from pathlib import Path

from googleapiclient.discovery import build
from mcp.server.fastmcp import FastMCP

from .auth import get_credentials

DEFAULT_DEST = Path(
    os.environ.get("GMAIL_ATT_DOWNLOAD_DIR", Path.home() / "Downloads" / "gmail-attachments")
)
TEXT_TYPES = ("text/", "application/json", "application/xml", "application/csv")

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

REMOTE = os.environ.get("GMAIL_ATT_REMOTE") == "1"

mcp = FastMCP(
    "gmail-attachments",
    host="0.0.0.0" if REMOTE else "127.0.0.1",
    port=int(os.environ.get("PORT", "8000")),
    stateless_http=True,
)


def _service():
    return build("gmail", "v1", credentials=get_credentials(), cache_discovery=False)


def _walk(part: dict):
    """Yield every part that is a real attachment (has a filename)."""
    if part.get("filename"):
        yield part
    for sub in part.get("parts", []) or []:
        yield from _walk(sub)


def _attachments(svc, message_id: str) -> list[dict]:
    msg = svc.users().messages().get(userId="me", id=message_id, format="full").execute()
    return list(_walk(msg["payload"]))


def _pick(parts: list[dict], filename: str, index: int) -> dict:
    matches = [p for p in parts if p["filename"] == filename]
    if not matches:
        names = ", ".join(p["filename"] for p in parts) or "none"
        raise ValueError(f"No attachment named {filename!r}. Available: {names}")
    if index >= len(matches):
        raise ValueError(f"{len(matches)} attachment(s) named {filename!r}; index {index} out of range")
    return matches[index]


def _bytes(svc, message_id: str, part: dict) -> bytes:
    body = part["body"]
    if "data" in body:
        data = body["data"]
    else:
        data = (
            svc.users().messages().attachments()
            .get(userId="me", messageId=message_id, id=body["attachmentId"])
            .execute()["data"]
        )
    return base64.urlsafe_b64decode(data + "=" * (-len(data) % 4))


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _xlsx_text(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                out.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


def _pptx_text(raw: bytes) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(io.BytesIO(raw)).slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return "\n".join(out)


def _safe_name(name: str) -> str:
    name = re.sub(r"[^\w.\- ]", "_", Path(name).name).strip(". ")
    return name or "attachment"


@mcp.tool()
def list_attachments(message_id: str) -> str:
    """List attachments on a Gmail message (name, type, size in bytes).

    message_id is the Gmail message id, e.g. from the Gmail connector's search_threads.
    """
    parts = _attachments(_service(), message_id)
    if not parts:
        return "No attachments."
    return "\n".join(
        f"{p['filename']} | {p.get('mimeType', '?')} | {p['body'].get('size', 0)} bytes"
        for p in parts
    )


def download_attachment(
    message_id: str, filename: str, dest_dir: str = "", index: int = 0
) -> str:
    """Save one attachment to disk and return the file path.

    filename: exact name from list_attachments. index: pick among duplicate names.
    dest_dir: folder to save into (default ~/Downloads/gmail-attachments).
    Existing files are never overwritten; a numeric suffix is added.
    """
    svc = _service()
    part = _pick(_attachments(svc, message_id), filename, index)
    folder = Path(dest_dir).expanduser() if dest_dir else DEFAULT_DEST
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / _safe_name(part["filename"])
    stem, suffix, n = target.stem, target.suffix, 1
    while target.exists():
        target = folder / f"{stem}-{n}{suffix}"
        n += 1
    target.write_bytes(_bytes(svc, message_id, part))
    return str(target)


if not REMOTE:  # hosted server has no persistent disk
    mcp.tool()(download_attachment)


@mcp.tool()
def read_attachment_text(
    message_id: str, filename: str, index: int = 0, max_chars: int = 20000
) -> str:
    """Return the text of an attachment without saving it.

    Supports PDF, docx, xlsx, pptx, and text files (txt, csv, json, xml, html, md).
    Other types (images, zip, doc, xls, ppt): use download_attachment (local server only).
    """
    svc = _service()
    part = _pick(_attachments(svc, message_id), filename, index)
    raw = _bytes(svc, message_id, part)
    mime = part.get("mimeType", "")
    if mime == "application/pdf" or filename.lower().endswith(".pdf"):
        from pypdf import PdfReader

        text = "\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
        if not text.strip():
            return "PDF has no extractable text (scanned?). Use download_attachment and view it."
    elif mime == DOCX or filename.lower().endswith(".docx"):
        text = _docx_text(raw)
    elif mime == XLSX or filename.lower().endswith(".xlsx"):
        text = _xlsx_text(raw)
    elif mime == PPTX or filename.lower().endswith(".pptx"):
        text = _pptx_text(raw)
    elif mime.startswith(TEXT_TYPES):
        text = raw.decode("utf-8", errors="replace")
    else:
        return f"Cannot extract text from {mime}. Use download_attachment."
    if len(text) > max_chars:
        return text[:max_chars] + f"\n\n[truncated, {len(text)} chars total]"
    return text


def main() -> None:
    mcp.run()


if __name__ == "__main__":
    main()
